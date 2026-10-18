"""Create a synthetic PPTX without starting Office or overwriting any output."""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path

TITLE = "AEGIS_SYNTHETIC_SMOKE_20261017"
BODY = "Synthetic validation content only."


def create(output_dir: Path) -> Path:
    """Write a fresh synthetic source and provenance; legacy PPT conversion is separate."""
    from pptx import Presentation
    output_dir.mkdir(parents=True, exist_ok=False)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = TITLE
    slide.placeholders[1].text = BODY
    target = output_dir / "synthetic-smoke.pptx"
    deck.save(target)
    metadata = dict(schema=1, synthetic=True, filename=target.name,
                    sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
                    historical_binary_reproduced=False, office_started=False,
                    legacy_ppt_conversion="not_run", legacy_read_contract="not_verified")
    (output_dir / "fixture.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    return target


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", required=True, type=Path)
    print(create(parser.parse_args().output_dir))
