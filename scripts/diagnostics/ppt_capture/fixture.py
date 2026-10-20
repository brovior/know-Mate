"""Create/verify fresh fixed-content synthetic PPTX files without Office."""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree

from .bundle import reject_links

TITLE = "AEGIS_CAPTURE_SYNTHETIC_V1"
BODY = "Synthetic capture validation only."
XML_TEXT = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
TEXT = TITLE + "\n" + BODY
TEXT_SHA256 = hashlib.sha256(TEXT.encode("utf-8")).hexdigest()


def create_fixture(output: Path) -> Path:
    """Generate one new fixed-content deck and its package-independent provenance."""
    from pptx import Presentation
    reject_links(output)
    output.mkdir(exist_ok=False)
    fixture_id = uuid4().hex
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = TITLE
    slide.placeholders[1].text = BODY
    target = output / ("capture-" + fixture_id + ".pptx")
    buffer = io.BytesIO()
    deck.save(buffer)
    # python-pptx's default template carries a binary printer profile. Remove
    # that unrelated template payload and its relationships from this new deck.
    with ZipFile(buffer) as source, ZipFile(target, "x", compression=ZIP_DEFLATED) as destination:
        relationship_ids: dict[str, set[str]] = {}
        replacements: dict[str, bytes] = {}
        for name in source.namelist():
            if name.endswith(".rels"):
                tree = ElementTree.fromstring(source.read(name))
                removed = {node.attrib["Id"] for node in tree if node.attrib.get("Type", "").endswith("/printerSettings")}
                if removed:
                    for node in list(tree):
                        if node.attrib.get("Id") in removed:
                            tree.remove(node)
                    parent = Path(name).parent.parent / Path(name).name.removesuffix(".rels")
                    relationship_ids[parent.as_posix()] = removed
                    replacements[name] = ElementTree.tostring(tree, encoding="utf-8", xml_declaration=True)
        for item in source.infolist():
            if item.filename.startswith("ppt/printerSettings/"):
                continue
            data = replacements.get(item.filename, source.read(item))
            if item.filename == "[Content_Types].xml" or item.filename in relationship_ids:
                tree = ElementTree.fromstring(data)
                if item.filename == "[Content_Types].xml":
                    for node in list(tree):
                        if node.attrib.get("PartName", "").startswith("/ppt/printerSettings/") or node.attrib.get("Extension") == "bin":
                            tree.remove(node)
                else:
                    for parent in tree.iter():
                        for child in list(parent):
                            if child.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id") in relationship_ids[item.filename]:
                                parent.remove(child)
                data = ElementTree.tostring(tree, encoding="utf-8", xml_declaration=True)
            destination.writestr(item, data)
    metadata = dict(schema=1, generator="ppt_capture.fixture.v1", fixture_id=fixture_id,
                    synthetic=True, office_started=False, filename=target.name,
                    sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
                    text_length=len(TEXT), text_sha256=TEXT_SHA256)
    with (output / "fixture.json").open("x", encoding="utf-8") as stream:
        json.dump(metadata, stream, indent=2)
        stream.write("\n")
    return target


def verify_fixture(path: Path, pins: dict[str, Any]) -> None:
    """Reject arbitrary documents, external relationships and altered synthetic content."""
    from .gate import digest, read_json
    metadata = read_json(path.parent / "fixture.json")
    fixture_id = metadata.get("fixture_id")
    if type(fixture_id) is not str or UUID(fixture_id).hex != fixture_id:
        raise ValueError("capture_fixture_id")
    if (metadata.get("schema") != 1 or metadata.get("generator") != "ppt_capture.fixture.v1"
            or metadata.get("synthetic") is not True or metadata.get("office_started") is not False
            or path.name != "capture-" + fixture_id + ".pptx" or metadata.get("filename") != path.name
            or metadata.get("sha256") != digest(path)
            or metadata.get("text_length") != len(TEXT) or metadata.get("text_sha256") != TEXT_SHA256
            or pins.get("fixture_text_length") != len(TEXT) or pins.get("fixture_text_sha256") != TEXT_SHA256
            or digest(path.parent / "fixture.json") != pins.get("fixture_provenance_sha256")):
        raise ValueError("capture_fixture_provenance")
    with ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or len(names) > 128:
            raise ValueError("capture_fixture_archive")
        slides = [name for name in names if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
        if slides != ["ppt/slides/slide1.xml"]:
            raise ValueError("capture_fixture_slides")
        for info in archive.infolist():
            if info.file_size > 1024 * 1024 or info.filename.endswith((".bin", ".vba")) or ".." in Path(info.filename).parts:
                raise ValueError("capture_fixture_member")
            if info.filename.endswith(".rels"):
                tree = ElementTree.fromstring(archive.read(info))
                if any(node.attrib.get("TargetMode") == "External" for node in tree):
                    raise ValueError("capture_fixture_external_relationship")
        texts = [node.text for node in ElementTree.fromstring(archive.read(slides[0])).iter(XML_TEXT)]
        if texts != [TITLE, BODY]:
            raise ValueError("capture_fixture_content")
