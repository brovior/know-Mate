"""Generate and verify the entire fixed synthetic OPC package, without Office."""
from __future__ import annotations

import hashlib
import io
import stat
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from scripts.diagnostics.ppt_capture.bundle import reject_links
from scripts.diagnostics.ppt_capture.fixture import BODY, TEXT, TEXT_SHA256, TITLE


def template_path() -> Path:
    """Return the actual python-pptx template used by this generator."""
    import pptx
    return Path(pptx.__file__).parent / "templates" / "default.pptx"


def template_sha256() -> str:
    """Pin the installed template rather than trusting a package version label."""
    from scripts.diagnostics.ppt_capture.gate import digest
    return digest(template_path())


def generator_sha256() -> str:
    """Pin the fixed generator source and its python-pptx dependency version."""
    import pptx
    from scripts.diagnostics.ppt_capture.gate import digest
    material = digest(Path(__file__)).encode("ascii") + b"\0" + pptx.__version__.encode("ascii")
    return hashlib.sha256(material).hexdigest()


def canonical_members() -> dict[str, bytes]:
    """Build a complete fixed member map in memory, excluding printer payloads."""
    from pptx import Presentation
    deck = Presentation(str(template_path()))
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = TITLE
    slide.placeholders[1].text = BODY
    stream = io.BytesIO()
    deck.save(stream)
    with ZipFile(stream) as source:
        members = {name: source.read(name) for name in source.namelist()
                   if not name.startswith("ppt/printerSettings/")}
    relationship_ids: dict[str, set[str]] = {}
    for name, data in list(members.items()):
        if not name.endswith(".rels"):
            continue
        tree = ET.fromstring(data)
        removed = {node.attrib["Id"] for node in tree
                   if node.attrib.get("Type", "").endswith("/printerSettings")}
        if removed:
            for node in list(tree):
                if node.attrib.get("Id") in removed:
                    tree.remove(node)
            parent = Path(name).parent.parent / Path(name).name.removesuffix(".rels")
            relationship_ids[parent.as_posix()] = removed
            members[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    for name, data in list(members.items()):
        if name != "[Content_Types].xml" and name not in relationship_ids:
            continue
        tree = ET.fromstring(data)
        for parent in tree.iter():
            for node in list(parent):
                printer_type = (name == "[Content_Types].xml" and
                                (node.attrib.get("PartName", "").startswith("/ppt/printerSettings/")
                                 or node.attrib.get("Extension") == "bin"))
                printer_reference = node.attrib.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
                ) in relationship_ids.get(name, set())
                if printer_type or printer_reference:
                    parent.remove(node)
        members[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    return members


@dataclass(frozen=True)
class Fixture:
    """One supervisor-created input; no caller-supplied document is accepted."""
    path: Path
    sha256: str
    generator_sha256: str
    template_sha256: str


def create_fixture(directory: Path) -> Fixture:
    """Create a fresh fixed deck and bind it to generator and template bytes."""
    reject_links(directory)
    directory.mkdir(mode=0o700, exist_ok=False)
    target = directory / "synthetic.pptx"
    with ZipFile(target, "x", compression=ZIP_DEFLATED) as archive:
        for name, data in canonical_members().items():
            archive.writestr(name, data)
    return Fixture(target, hashlib.sha256(target.read_bytes()).hexdigest(),
                   generator_sha256(), template_sha256())


def verify_fixture(fixture: Fixture) -> None:
    """Reject any changed member, including notes, customXml and document properties."""
    reject_links(fixture.path)
    if (fixture.path.name != "synthetic.pptx" or not fixture.path.is_file()
            or fixture.generator_sha256 != generator_sha256()
            or fixture.template_sha256 != template_sha256()
            or fixture.path.stat().st_size > 4 * 1024 * 1024):
        raise ValueError("lifetime_fixture_pin")
    expected = canonical_members()
    with ZipFile(fixture.path) as archive:
        infos = archive.infolist()
        names = [info.filename for info in infos]
        if len(names) != len(set(names)) or set(names) != set(expected):
            raise ValueError("lifetime_fixture_member_set")
        for info in infos:
            if (info.file_size != len(expected[info.filename]) or info.flag_bits & 1
                    or stat.S_ISLNK(info.external_attr >> 16)
                    or "\\" in info.filename or ".." in Path(info.filename).parts):
                raise ValueError("lifetime_fixture_member_type_or_size")
            if archive.read(info) != expected[info.filename]:
                raise ValueError("lifetime_fixture_content")
    if hashlib.sha256(fixture.path.read_bytes()).hexdigest() != fixture.sha256:
        raise ValueError("lifetime_fixture_digest")
