"""업무 데이터 없이 표·병합·줄바꿈·중첩 그룹을 가진 합성 PPT를 만든다."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Inches


def create_pptx(path: Path, rows: int, columns: int, slides: int) -> None:
    """속도 비교용 합성 문서를 지정한 실험 디렉터리에 저장한다."""
    prs = Presentation()
    for number in range(slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(0.1), Inches(0.1), Inches(8), Inches(0.4))
        title.text = f"합성 검증 슬라이드 {number + 1}"
        table = slide.shapes.add_table(rows, columns, Inches(0.1), Inches(0.7), Inches(8), Inches(4)).table
        for row in range(rows):
            for column in range(columns):
                table.cell(row, column).text = "" if row == rows - 1 else f"값 {number}/{row}/{column}"
        table.cell(0, 0).text = "첫 줄\n둘째 줄"
        if rows > 2 and columns > 1:
            table.cell(1, 0).merge(table.cell(1, 1))
        outer = slide.shapes.add_group_shape()
        box = outer.shapes.add_textbox(Inches(0.1), Inches(5), Inches(3), Inches(0.4))
        box.text = "그룹 첫 도형"
        inner = outer.shapes.add_group_shape()
        box = inner.shapes.add_textbox(Inches(3), Inches(5), Inches(3), Inches(0.4))
        box.text = "중첩 그룹 도형"
    prs.slides.add_slide(prs.slide_layouts[6])  # 빈 슬라이드의 위치 보존도 확인
    prs.save(str(path))


def create_legacy(app: Any, source: Path, destination: Path) -> None:
    """직접 만든 PPTX를 구형 PPT로 저장해 변환 실험의 입력을 준비한다."""
    from scripts.ppt_probe.readers import Snapshot, opened

    def save(prs: Any) -> Snapshot:
        prs.SaveCopyAs(str(destination.resolve()), 1)  # ppSaveAsPresentation (.ppt)
        return Snapshot()

    opened(app, source, save)
