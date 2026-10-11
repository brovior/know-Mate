"""PowerPoint COM·파일·내부 추가 기능의 결과를 메모리에서 비교한다."""
from __future__ import annotations

import json
import math
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from knowmate.secure import com_reader
from knowmate.secure.plain_reader import PlainReader
from knowmate.secure.ppt_diagnostics import PptReadDiagnostics
from knowmate.secure.text_util import format_table

PROG_ID = "AegisDesk.PptReadProbe"
MAX_REPLY_CHARS = 32 * 1024 * 1024


class ProbeFailure(RuntimeError):
    """본문을 포함하지 않는 검증 실패 코드."""


class ProbeCloseFailure(ProbeFailure):
    """문서 닫기 미확인 상태에서는 후속 COM 실험을 중단한다."""

    def __init__(self, exc: Exception) -> None:
        self.hresult = com_reader.hresult_of(exc)
        super().__init__("document_close_unconfirmed")


@dataclass
class Snapshot:
    """출력하지 않는 추출 내용과 수치 진단 정보를 보관한다."""

    slides: list[list[str]] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=lambda: {
        "shapes": 0, "groups": 0, "tables": 0, "cells": 0, "errors": 0,
    })
    seconds: dict[str, float] = field(default_factory=dict)

    @property
    def text(self) -> str:
        """현재 리더와 같은 슬라이드 구분으로 내용을 조합한다."""
        return "\n\n".join("\n".join(s) for s in self.slides if s)


def compare(reference: Snapshot, candidate: Snapshot) -> dict[str, Any]:
    """내용을 출력하지 않고 슬라이드·순서·건수의 일치 여부를 반환한다."""
    mismatch = [i + 1 for i in range(max(len(reference.slides), len(candidate.slides)))
                if i >= len(reference.slides) or i >= len(candidate.slides)
                or reference.slides[i] != candidate.slides[i]]
    counts_equal = reference.counts == candidate.counts
    return {
        "equal": not mismatch and counts_equal and reference.counts["errors"] == 0,
        "text_equal": reference.text == candidate.text,
        "counts_equal": counts_equal, "mismatch_slides": mismatch,
        "characters": len(candidate.text), "slides": len(candidate.slides),
        "counts": candidate.counts, "seconds": candidate.seconds,
    }


def _count_plain(shape: Any, counts: dict[str, int]) -> None:
    counts["shapes"] += 1
    if int(shape.shape_type) == 6:
        counts["groups"] += 1
        for child in shape.shapes:
            _count_plain(child, counts)
    elif shape.has_table:
        counts["tables"] += 1
        counts["cells"] += len(shape.table.rows) * len(shape.table.columns)


def read_plain(path: Path) -> Snapshot:
    """기존 PPTX 파서와 같은 형식으로 파일을 직접 읽는다."""
    from pptx import Presentation

    result = Snapshot()
    started = time.perf_counter()
    prs = Presentation(str(path))
    reader = PlainReader()
    for slide in prs.slides:
        texts: list[str] = []
        for shape in slide.shapes:
            _count_plain(shape, result.counts)
            texts.extend(reader._iter_shape_texts(shape))
        result.slides.append([t for t in texts if t.strip()])
    result.seconds["parse"] = time.perf_counter() - started
    return result


def read_baseline(prs: Any, path: Path) -> Snapshot:
    """현재 COM 도형 리더를 그대로 호출하고 오류·건수를 함께 확인한다."""
    result = Snapshot()
    with PptReadDiagnostics(str(path)) as diagnostics:
        collection = diagnostics.read("slides", lambda: prs.Slides, 0)
        for index, slide in diagnostics.items(collection, "slides", 0):
            shapes = diagnostics.read("shapes", lambda: slide.Shapes, index)
            texts: list[str] = []
            for position, shape in diagnostics.items(shapes, "shapes", index):
                texts.extend(com_reader._ppt_shape_texts(shape, diagnostics, index, (position,)))
            result.slides.append([t for t in texts if t.strip()])
            diagnostics.slide_finished(index)
        for stats in diagnostics._stats.values():
            for name in result.counts:
                result.counts[name] += getattr(stats, name)
    return result


def decode_slide(reply: str, slide: int) -> tuple[list[str], dict[str, int], float]:
    """추가 기능의 응답을 검증하고 기존 표 정렬을 적용한다."""
    if not isinstance(reply, str) or len(reply) > MAX_REPLY_CHARS:
        raise ProbeFailure("invalid_reply_size")
    payload = json.loads(reply)
    if not isinstance(payload, dict) or payload.get("schema") != 1 or payload.get("slide") != slide:
        raise ProbeFailure("invalid_reply_schema")
    counts = payload.get("counts")
    keys = {"shapes", "groups", "tables", "cells", "errors"}
    if not isinstance(counts, dict) or set(counts) != keys:
        raise ProbeFailure("invalid_reply_counts")
    if any(type(v) is not int or v < 0 for v in counts.values()):
        raise ProbeFailure("invalid_reply_counts")
    seconds = payload.get("seconds")
    if type(seconds) not in (int, float) or not math.isfinite(seconds) or seconds < 0:
        raise ProbeFailure("invalid_reply_time")
    nodes = payload.get("nodes")
    if not isinstance(nodes, list):
        raise ProbeFailure("invalid_reply_nodes")
    texts: list[str] = []
    for node in nodes:
        if not isinstance(node, dict):
            raise ProbeFailure("invalid_reply_node")
        if node.get("kind") == "text" and isinstance(node.get("text"), str):
            text = node["text"].strip()
        elif node.get("kind") == "table" and isinstance(node.get("cells"), list):
            grid = node["cells"]
            if any(not isinstance(row, list) or any(not isinstance(c, str) for c in row) for row in grid):
                raise ProbeFailure("invalid_reply_cells")
            text = format_table(grid)
        else:
            raise ProbeFailure("invalid_reply_node")
        if text.strip():
            texts.append(text)
    return texts, counts, float(seconds)


def read_addin(prs: Any, path: Path, service: Any) -> Snapshot:
    """슬라이드당 한 번 내부 추출을 요청하고 결과는 메모리에만 보관한다."""
    result = Snapshot()
    internal = 0.0
    with PptReadDiagnostics(str(path)) as diagnostics:
        count = int(prs.Slides.Count)
        for index in range(1, count + 1):
            reply = diagnostics.read("addin_slide", lambda: service.ExtractSlide(prs, index), index)
            texts, counts, seconds = decode_slide(reply, index)
            result.slides.append(texts)
            internal += seconds
            for name in result.counts:
                result.counts[name] += counts[name]
    result.seconds["internal_read"] = internal
    return result


def opened(app: Any, path: Path, reader: Callable[[Any], Snapshot]) -> Snapshot:
    """읽기 전용으로 열고 닫는 시간을 포함해 방법별 전체 시간을 잰다."""
    started = time.perf_counter()
    prs = None
    result = None
    fatal = False
    try:
        prs = app.Presentations.Open(str(path.resolve()), ReadOnly=True, WithWindow=False)
        if prs is None:
            raise com_reader.OfficeComPoisonError("POWERPNT.EXE")
        open_seconds = time.perf_counter() - started
        read_started = time.perf_counter()
        result = reader(prs)
        result.seconds.update(open=open_seconds, read=time.perf_counter() - read_started)
    except Exception as exc:
        fatal = isinstance(exc, com_reader.OfficeComPoisonError) or com_reader.is_fatal_transport_hresult(
            com_reader.hresult_of(exc))
        raise
    finally:
        if prs is not None and not fatal:
            close_started = time.perf_counter()
            try:
                prs.Close()
            except Exception as exc:
                raise ProbeCloseFailure(exc) from exc
            if result is not None:
                result.seconds["close"] = time.perf_counter() - close_started
    assert result is not None
    result.seconds["total"] = time.perf_counter() - started
    return result


def convert_synthetic(app: Any, source: Path, destination: Path) -> Snapshot:
    """합성 PPT만 PPTX로 저장한 뒤 직접 파싱하며 변환 비용도 포함한다."""
    def convert(prs: Any) -> Snapshot:
        started = time.perf_counter()
        prs.SaveCopyAs(str(destination.resolve()), 24)  # ppSaveAsOpenXMLPresentation
        saved = time.perf_counter() - started
        result = read_plain(destination)
        result.seconds["save"] = saved
        return result

    return opened(app, source, convert)
